from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


SCRIPT_PATH = Path(__file__).with_name("generate-client-template-manifest.py")
SPEC = importlib.util.spec_from_file_location("generate_client_template_manifest", SCRIPT_PATH)
assert SPEC and SPEC.loader
GENERATOR = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(GENERATOR)


class GenerateClientTemplateManifestTest(unittest.TestCase):
    def test_normalizes_all_supported_boolean_markers(self) -> None:
        self.assertIs(GENERATOR.parse_boolean("SI"), True)
        self.assertIs(GENERATOR.parse_boolean("S"), True)
        self.assertIs(GENERATOR.parse_boolean("NO"), False)
        self.assertIs(GENERATOR.parse_boolean("N"), False)

    def test_pim2_variants_taxonomy_normalization_and_relations(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            workbook_path = Path(temporary_directory) / "PLANTILLAS PIM(2).xlsx"
            self._make_workbook(workbook_path)

            manifest = GENERATOR.build_manifest(workbook_path)

        self.assertEqual(manifest["version"], 2)
        self.assertEqual(manifest["source"]["sheetCount"], 9)
        self.assertEqual(manifest["source"]["templateCount"], 2)
        self.assertEqual(manifest["source"]["productCount"], 2)
        self.assertEqual(len(manifest["categories"]), 6)

        category = next(item for item in manifest["categories"] if item["categoryNumber"] == "22")
        root = next(item for item in manifest["categories"] if item["categoryPath"] == ["FILTROS"])
        middle = next(
            item
            for item in manifest["categories"]
            if item["categoryPath"] == ["FILTROS", "FILTRO DE AIRE"]
        )
        self.assertEqual((root["level"], root["parentSlug"]), (1, None))
        self.assertEqual((middle["level"], middle["parentSlug"]), (2, root["slug"]))
        self.assertNotIn("templateVariants", root)
        self.assertNotIn("templateVariants", middle)
        self.assertEqual((category["level"], category["parentSlug"]), (3, middle["slug"]))
        self.assertEqual(category["categoryPath"], ["FILTROS", "FILTRO DE AIRE", "FILTRO DE AIRE"])
        self.assertEqual(
            [variant["variantKey"] for variant in category["templateVariants"]],
            ["22.1", "22.2"],
        )
        self.assertEqual(len({template["categorySlug"] for template in manifest["templates"]}), 1)
        self.assertEqual(
            {template["slug"] for template in manifest["templates"]},
            {
                f"{category['slug']}--22-1",
                f"{category['slug']}--22-2",
            },
        )

        ecommerce = next(
            definition for definition in manifest["definitions"] if definition["key"] == "ecommerce"
        )
        self.assertEqual(ecommerce["dataType"], "boolean")
        self.assertEqual(
            [template["products"][0]["values"]["ecommerce"] for template in manifest["templates"]],
            [True, False],
        )
        self.assertNotIn("", {definition["key"] for definition in manifest["definitions"]})
        self.assertNotIn(
            "campo_con_error",
            manifest["templates"][0]["products"][0]["values"],
        )

        application = manifest["relations"]["applications"][0]
        self.assertEqual(
            {key: application[key] for key in ("tipo", "combustible", "traccion", "cilindraje")},
            {"tipo": "HB", "combustible": "GSL", "traccion": "4X2", "cilindraje": "1.6L"},
        )
        self.assertEqual(manifest["relations"]["oem"][0]["codigo_oem"], "OEM-1")
        self.assertEqual(manifest["relations"]["unifiedCodes"][0]["codigo_unificador"], "U-1")
        diagnostic_codes = {diagnostic["code"] for diagnostic in manifest["diagnostics"]}
        self.assertIn("BROKEN_MARKER_IGNORED", diagnostic_codes)
        self.assertIn("UNLABELED_COLUMNS_IGNORED", diagnostic_codes)

    @staticmethod
    def _make_workbook(path: Path) -> None:
        workbook = Workbook()
        categories = workbook.active
        categories.title = "CATEGORIAS-BUSQUEDA"
        categories.append(
            [
                "Categoria Principal",
                "Categoria Segundo Nivel",
                "Categoria Tercer Nivel",
                "APLICACIÓN",
                "TECDOC",
                "FABRICANTE",
                "ARCHIVO",
                "MANUAL",
                "Nro. Plantilla",
                "OBSERVACIÓN",
            ]
        )
        categories.append(
            ["FILTROS", "FILTRO DE AIRE", "FILTRO DE AIRE", "AUTOMOTRIZ", 1, 2, 3, 4, 22, None]
        )
        categories.append(
            ["MOTOR", "REPUESTOS", "PENDIENTE", None, None, None, None, None, None, None]
        )

        for title in (
            "REPORTE ATRIBUTOS",
            "REPORTE-GENERAL(RESULTADOS)",
            "REPORTE-RESULTADOS DE PERIODO",
        ):
            workbook.create_sheet(title).cell(1, 1, "No es una plantilla")

        GenerateClientTemplateManifestTest._add_template(
            workbook, "22.1 FILTRO RECTANGULAR", "SKU-1", "U-1", "SI", "#VALUE!"
        )
        GenerateClientTemplateManifestTest._add_template(
            workbook, "22.2 FILTRO CILINDRICO", "SKU-2", "U-2", "N", "válido"
        )

        applications = workbook.create_sheet("APLICACIONES")
        GenerateClientTemplateManifestTest._set_relation_headers(
            applications,
            [
                "Código Unificador",
                "Marca Vehiculo",
                "Modelo",
                "Tipo",
                "Combustible",
                "Tracción",
                "Cilindraje",
                "Año Desde",
                "Año Hasta",
                "Area",
                "Subarea",
                "Estado",
                "SKU - Asociados",
                "Fuente",
            ],
            ["U-1", "SUZUKI", "SX4", "HB", "GSL", "4X2", "1.6L", 2013, 2019, "RUEDA", "POSTERIOR", "OK", "SKU-1", "TECDOC"],
        )
        oem = workbook.create_sheet("OEM")
        GenerateClientTemplateManifestTest._set_relation_headers(
            oem,
            ["Código Unificador", "Códigp OEM", "Marca", "Estado", "SKU - Asociados", "Fuente"],
            ["U-1", "OEM-1", "SUZUKI", "OK", "SKU-1", "TECDOC"],
        )
        unified = workbook.create_sheet("CODIGO UNIFICADOR")
        GenerateClientTemplateManifestTest._set_relation_headers(
            unified,
            ["Código Unificador", "Código de Articulo - CDR", "Código de Proveedor", "Marca"],
            ["U-1", "SKU-1", "PROV-1", "MARCA"],
        )
        workbook.save(path)

    @staticmethod
    def _add_template(
        workbook: Workbook,
        title: str,
        sku: str,
        unified_code: str,
        ecommerce: str,
        error_value: str,
    ) -> None:
        sheet = workbook.create_sheet(title)
        headers = [
            "Código de Articulo - CDR",
            "Código de Proveedor",
            "Código Unificador",
            "Categoria Principal",
            "Categoria Segundo Nivel",
            "Categoria Tercer Nivel",
            "Tipo de Aplicación",
            "Marca de Artículo",
            "Ecommerce",
            "Campo con error",
        ]
        for column, header in enumerate(headers, 3):
            sheet.cell(1, column, "ALFANUMERICO")
            sheet.cell(2, column, "OBLIGATORIO" if column <= 11 else "OPCIONAL")
            sheet.cell(3, column, "DESDE ERP" if column <= 11 else "USUARIO")
            sheet.cell(4, column, "S")
            sheet.cell(5, column, "SI")
            sheet.cell(7, column, "SI")
            sheet.cell(8, column, "S")
            sheet.cell(9, column, "SI")
            sheet.cell(10, column, header)
        values = [
            sku,
            f"PROV-{sku}",
            unified_code,
            "FILTROS",
            "FILTRO DE AIRE",
            "FILTRO DE AIRE",
            "AUTOMOTRIZ",
            "MARCA",
            ecommerce,
            error_value,
        ]
        for column, value in enumerate(values, 3):
            sheet.cell(11, column, value)
        sheet.cell(11, 14, "debe ignorarse")

    @staticmethod
    def _set_relation_headers(sheet: object, headers: list[str], values: list[object]) -> None:
        for column, header in enumerate(headers, 3):
            sheet.cell(10, column, header)
        for column, value in enumerate(values, 3):
            sheet.cell(11, column, value)


if __name__ == "__main__":
    unittest.main()
