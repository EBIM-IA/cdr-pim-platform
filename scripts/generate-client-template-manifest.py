#!/usr/bin/env python3
"""Generate the source-traceable local catalogue manifest from a PIM workbook.

Usage:
    python3 scripts/generate-client-template-manifest.py INPUT.xlsx [OUTPUT.json]

The JSON is deterministic runtime input for local/test data. Spreadsheet parsing deliberately
stays outside the API runtime. Source anomalies which can be handled without inventing data are
preserved as structured diagnostics instead of being silently repaired.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SEARCH_SHEETS = ("CATEGORIAS-BUSQUEDA", "BUSQUEDA")
NON_TEMPLATE_SHEETS = {
    *SEARCH_SHEETS,
    "REPORTE ATRIBUTOS",
    "REPORTE-GENERAL(RESULTADOS)",
    "REPORTE-RESULTADOS DE PERIODO",
    "CODIGO UNIFICADOR",
    "OEM",
    "HOMOLOGOS",
    "APLICACIONES",
}
ASSET_CHARACTERS = {"ARCHIVO", "IMAGEN", "ARCHIVO / IMAGEN", "ARCHIVO/IMAGEN"}
BASE_KEYS = {
    "codigo articulo": "codigo_articulo",
    "codigo de articulo": "codigo_articulo",
    "codigo articulo cdr": "codigo_articulo",
    "codigo de articulo cdr": "codigo_articulo",
    "codigo proveedor": "codigo_proveedor",
    "codigo de proveedor": "codigo_proveedor",
    "codigo unificador": "codigo_unificador",
    "categoria principal": "categoria_principal",
    "categoria segundo nivel": "categoria_segundo_nivel",
    "categoria tercer nivel": "categoria_tercer_nivel",
    "linea": "linea",
    "codigo de linea": "linea",
    "tipo de aplicacion": "tipo_aplicacion",
    "aplicacion": "aplicacion",
    "marca": "marca",
    "marca de articulo": "marca",
    "marca vehiculo": "marca_vehiculo",
    "modelo": "modelo",
    "tipo": "tipo",
    "combustible": "combustible",
    "traccion": "traccion",
    "cilindraje": "cilindraje",
    "ano desde": "ano_desde",
    "ano hasta": "ano_hasta",
    "area": "area",
    "subarea": "subarea",
    "estado": "estado",
    "estado sku": "estado_sku",
    "sku asociados": "sku_asociados",
    "fuente": "fuente",
    "codigo oem": "codigo_oem",
    "codigp oem": "codigo_oem",
    "codigo homologo": "codigo_homologo",
    "marca homologo": "marca_homologo",
    "codigo de fabrica": "codigo_fabrica",
    "grupo": "grupo",
    "homologos": "homologos",
    "oem": "oem",
    "frecuencia": "frecuencia_articulo",
    "ecommerce": "ecommerce",
    "bloqueado venta": "bloqueado_venta",
    "bloqueado compra": "bloqueado_compra",
    "especificacion tecnica observacion": "especificacion_observacion",
    "especificacion observacion": "especificacion_observacion",
}
BASE_ERP_KEYS = {
    "codigo_articulo",
    "codigo_proveedor",
    "codigo_unificador",
    "categoria_principal",
    "categoria_segundo_nivel",
    "categoria_tercer_nivel",
    "linea",
    "tipo_aplicacion",
    "marca",
    "frecuencia_articulo",
    "ecommerce",
    "bloqueado_venta",
    "bloqueado_compra",
    "especificacion_observacion",
}
EMPTY_MARKERS = {"", "-", "none"}
BROKEN_MARKER = re.compile(r"^#(?:value!|ref!|div/0!|name\?|n/a|num!|null!)$", re.I)


def cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("_x000D_", "").strip()


def fold(value: object) -> str:
    normalized = unicodedata.normalize("NFKD", cell_text(value))
    without_marks = "".join(char for char in normalized if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", without_marks).strip().casefold()


def words(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", fold(value)).strip()


def slug(value: object, limit: int = 100) -> str:
    return re.sub(r"[^a-z0-9]+", "-", fold(value)).strip("-")[:limit]


def attribute_key(value: object) -> str:
    normalized = words(value)
    return BASE_KEYS.get(normalized, normalized.replace(" ", "_")[:100])


def yes(value: object) -> bool:
    return fold(value) in {"si", "s", "genera el sistema"}


def source_authority(value: object, key: str) -> str:
    if key in BASE_ERP_KEYS:
        return "erp"
    normalized = fold(value)
    if "erp" in normalized:
        return "erp"
    if "genera el sistema" in normalized:
        return "calculated"
    return "pim"


def is_broken_marker(value: object) -> bool:
    return bool(BROKEN_MARKER.fullmatch(cell_text(value)))


def meaningful(value: object) -> bool:
    return fold(value) not in EMPTY_MARKERS and not is_broken_marker(value)


def parse_number(value: object) -> float | int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    raw = cell_text(value).replace(" ", "").replace(",", ".")
    if not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", raw):
        return None
    parsed = float(raw)
    return int(parsed) if parsed.is_integer() else parsed


def parse_boolean(value: object) -> bool | None:
    if isinstance(value, bool):
        return value
    normalized = fold(value)
    if normalized in {"si", "s"}:
        return True
    if normalized in {"no", "n"}:
        return False
    return None


def normalize_template_number(value: object) -> str | None:
    if value is None or not meaningful(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    raw = cell_text(value).rstrip(".")
    match = re.fullmatch(r"(\d+)(?:\.(\d+))?", raw)
    if not match:
        return raw
    return match.group(1) + (f".{match.group(2)}" if match.group(2) else "")


def category_number(variant_key: str) -> str:
    return variant_key.split(".", 1)[0]


def explicit_unit(label: str) -> str | None:
    normalized = fold(label).replace("º", "°")
    candidates = [
        (r"(?:\[|\(|\s)mm(?:\]|\)|\s|:|$)", "mm"),
        (r"(?:\[|\(|\s)kg(?:\]|\)|\s|:|$)", "kg"),
        (r"(?:\[|\(|\s)nm(?:\]|\)|\s|:|$)", "Nm"),
        (r"(?:\[|\(|\s)°c(?:\]|\)|\s|:|$)", "°C"),
    ]
    return next((unit for pattern, unit in candidates if re.search(pattern, normalized)), None)


def asset_type(label: str) -> str | None:
    normalized = words(label)
    if "msds" in normalized or "seguridad" in normalized:
        return "MSDS"
    if "certificado" in normalized:
        return "CERT"
    if "fotografia" in normalized or "foto" in normalized:
        return "FOTO"
    if "plano" in normalized:
        return "PLANO"
    if "ficha tecnica" in normalized:
        return "FT"
    return None


def find_header_row(sheet: Any) -> int:
    for row in range(1, min(sheet.max_row, 20) + 1):
        if any(
            attribute_key(sheet.cell(row, column).value) == "codigo_articulo"
            for column in range(1, sheet.max_column + 1)
        ):
            return row
    raise ValueError(f"{sheet.title}: no se encontró la fila de encabezados")


def sheet_number(title: str) -> str | None:
    match = re.match(r"^\s*(\d+(?:\.\d+)?)\.?", title)
    return normalize_template_number(match.group(1)) if match else None


def sheet_name(title: str) -> str:
    return re.sub(r"^\s*\d+(?:\.\d+)?\.?\s*", "", title).strip()


class Diagnostics:
    def __init__(self) -> None:
        self.items: list[dict[str, Any]] = []

    def add(
        self,
        code: str,
        message: str,
        *,
        sheet: str | None = None,
        location: str | None = None,
        count: int | None = None,
        details: list[str] | None = None,
    ) -> None:
        item: dict[str, Any] = {"severity": "warning", "code": code, "message": message}
        if sheet:
            item["sheet"] = sheet
        if location:
            item["location"] = location
        if count is not None:
            item["count"] = count
        if details:
            item["details"] = details
        self.items.append(item)


def unique_slug(path: list[str], used: dict[str, list[str]]) -> str:
    candidate = slug(" ".join(path)) or "categoria"
    previous = used.get(candidate)
    if previous is None or previous == path:
        used[candidate] = path
        return candidate
    suffix = hashlib.sha1("/".join(path).encode("utf-8")).hexdigest()[:6]
    candidate = f"{candidate[:93]}-{suffix}"
    used[candidate] = path
    return candidate


def read_search_matrix(book: Any, diagnostics: Diagnostics) -> list[dict[str, Any]]:
    search_sheet_name = next((name for name in SEARCH_SHEETS if name in book.sheetnames), None)
    if not search_sheet_name:
        raise ValueError("No se encontró CATEGORIAS-BUSQUEDA ni BUSQUEDA")
    sheet = book[search_sheet_name]
    leaf_categories_by_path: dict[tuple[str, ...], dict[str, Any]] = {}
    for row in range(2, sheet.max_row + 1):
        if search_sheet_name == "CATEGORIAS-BUSQUEDA":
            levels = [cell_text(sheet.cell(row, column).value) for column in range(1, 4)]
            application_column, priority_start, number_column, observation_column = 4, 5, 9, 10
        else:
            # Compatibility with the first PIM workbook, which had a single category level.
            levels = [cell_text(sheet.cell(row, 2).value)]
            application_column, priority_start, number_column, observation_column = 3, 4, 1, 0
        if not any(levels):
            continue
        path = [level for level in levels if level]
        if len(path) != len(levels):
            diagnostics.add(
                "INCOMPLETE_CATEGORY_PATH",
                "La categoría no contiene todos los niveles declarados por la matriz.",
                sheet=sheet.title,
                location=f"fila {row}",
                details=levels,
            )
        key = tuple(fold(level) for level in path)
        existing = leaf_categories_by_path.get(key)
        if existing:
            existing["sourceRows"].append(row)
            diagnostics.add(
                "DUPLICATE_CATEGORY_PATH",
                "La ruta de categoría aparece más de una vez; se consolidó sin crear otra categoría.",
                sheet=sheet.title,
                location=f"filas {existing['excelRow']}, {row}",
                details=path,
            )
            continue
        priorities = {
            "tecdoc": sheet.cell(row, priority_start).value,
            "fabricante": sheet.cell(row, priority_start + 1).value,
            "archivo": sheet.cell(row, priority_start + 2).value,
            "manual": sheet.cell(row, priority_start + 3).value,
        }
        source_priority: dict[str, int] = {}
        invalid_priorities: list[str] = []
        used_priorities: set[int] = set()
        for priority_name, value in priorities.items():
            if not isinstance(value, (int, float)):
                continue
            parsed_priority = int(value)
            if parsed_priority not in {1, 2, 3, 4} or parsed_priority in used_priorities:
                invalid_priorities.append(f"{priority_name}={cell_text(value)}")
                continue
            source_priority[priority_name] = parsed_priority
            used_priorities.add(parsed_priority)
        if invalid_priorities:
            diagnostics.add(
                "INVALID_SOURCE_PRIORITY",
                "Se ignoraron prioridades de fuente fuera de 1-4 o repetidas.",
                sheet=sheet.title,
                location=f"fila {row}",
                details=invalid_priorities,
            )
        observation = (
            cell_text(sheet.cell(row, observation_column).value) if observation_column else ""
        )
        number = normalize_template_number(sheet.cell(row, number_column).value)
        record = {
            "excelRow": row,
            "sourceRows": [row],
            "categoryNumber": category_number(number) if number else None,
            "name": path[-1],
            "categoryPath": path,
            "application": cell_text(sheet.cell(row, application_column).value) or None,
            "sourcePriority": source_priority,
            "observation": observation or None,
        }
        leaf_categories_by_path[key] = record
        if observation:
            diagnostics.add(
                "SOURCE_OBSERVATION",
                "La matriz contiene una observación que requiere revisión funcional.",
                sheet=sheet.title,
                location=f"fila {row}",
                details=[observation],
            )
    categories_by_path: dict[tuple[str, ...], dict[str, Any]] = {}
    used_slugs: dict[str, list[str]] = {}
    for leaf in leaf_categories_by_path.values():
        parent_slug: str | None = None
        for level in range(1, len(leaf["categoryPath"]) + 1):
            path = leaf["categoryPath"][:level]
            key = tuple(fold(part) for part in path)
            category = categories_by_path.get(key)
            if category is None:
                category = (
                    leaf
                    | {
                        "sourceRows": list(leaf["sourceRows"]),
                        "name": path[-1],
                        "slug": unique_slug(path, used_slugs),
                        "level": level,
                        "parentSlug": parent_slug,
                        "categoryPath": path,
                    }
                    if level == len(leaf["categoryPath"])
                    else {
                        "excelRow": leaf["excelRow"],
                        "sourceRows": list(leaf["sourceRows"]),
                        "categoryNumber": None,
                        "name": path[-1],
                        "slug": unique_slug(path, used_slugs),
                        "level": level,
                        "parentSlug": parent_slug,
                        "categoryPath": path,
                        "application": None,
                        "sourcePriority": {},
                        "observation": None,
                    }
                )
                categories_by_path[key] = category
            else:
                category["sourceRows"] = sorted(
                    {*category["sourceRows"], *leaf["sourceRows"]}
                )
            parent_slug = category["slug"]
    return list(categories_by_path.values())


def relation_rows(sheet: Any, diagnostics: Diagnostics) -> list[dict[str, object]]:
    header_row = 10
    headers = [
        attribute_key(sheet.cell(header_row, column).value)
        for column in range(1, sheet.max_column + 1)
    ]
    records: list[dict[str, object]] = []
    broken_locations: list[str] = []
    for row in range(header_row + 1, sheet.max_row + 1):
        record: dict[str, object] = {}
        for column, key in enumerate(headers, 1):
            value = sheet.cell(row, column).value
            if is_broken_marker(value):
                broken_locations.append(sheet.cell(row, column).coordinate)
                continue
            if not key or not meaningful(value):
                continue
            parsed = parse_number(value) if key in {"ano_desde", "ano_hasta"} else None
            record[key] = parsed if parsed is not None else value
        if record:
            records.append(record)
    if broken_locations:
        diagnostics.add(
            "BROKEN_MARKER_IGNORED",
            "Se ignoraron marcadores de error de Excel en la relación.",
            sheet=sheet.title,
            count=len(broken_locations),
            details=broken_locations[:20],
        )
    return records


def build_manifest(source_path: Path) -> dict[str, Any]:
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    book = load_workbook(source_path, data_only=True, read_only=False)
    diagnostics = Diagnostics()
    categories = read_search_matrix(book, diagnostics)
    categories_by_number = {
        category["categoryNumber"]: category
        for category in categories
        if category["categoryNumber"] is not None
    }

    raw_templates: list[dict[str, Any]] = []
    occurrences: dict[str, list[dict[str, Any]]] = defaultdict(list)
    used_template_slugs: set[str] = set()

    for sheet in book.worksheets:
        if sheet.title in NON_TEMPLATE_SHEETS:
            continue
        variant_key = sheet_number(sheet.title)
        if not variant_key:
            diagnostics.add(
                "NON_TEMPLATE_SHEET_SKIPPED",
                "La hoja no tiene número de plantilla y fue excluida.",
                sheet=sheet.title,
            )
            continue
        base_number = category_number(variant_key)
        category = categories_by_number.get(base_number)
        if category is None:
            raise ValueError(
                f"{sheet.title}: no existe categoría para el número de plantilla {base_number}"
            )
        header_row = find_header_row(sheet)
        template_name = sheet_name(sheet.title)
        template_slug = f"{category['slug']}--{variant_key.replace('.', '-')}"
        if template_slug in used_template_slugs:
            raise ValueError(f"Slug de plantilla duplicado: {template_slug}")
        used_template_slugs.add(template_slug)
        if fold(template_name) != fold(category["name"]):
            diagnostics.add(
                "TEMPLATE_CATEGORY_NAME_DIFFERENCE",
                "El nombre de la hoja no coincide exactamente con la categoría mapeada; no se autocorrigió.",
                sheet=sheet.title,
                location=f"{category['slug']} (fila {category['excelRow']})",
                details=[template_name, category["name"]],
            )

        unlabeled_columns: list[str] = []
        for column in range(3, sheet.max_column + 1):
            if cell_text(sheet.cell(header_row, column).value):
                continue
            if any(
                meaningful(sheet.cell(row, column).value)
                for row in range(1, sheet.max_row + 1)
                if row != header_row
            ):
                unlabeled_columns.append(sheet.cell(header_row, column).column_letter)
        if unlabeled_columns:
            diagnostics.add(
                "UNLABELED_COLUMNS_IGNORED",
                "Se ignoraron columnas con contenido pero sin encabezado.",
                sheet=sheet.title,
                count=len(unlabeled_columns),
                details=unlabeled_columns,
            )

        attributes: list[dict[str, Any]] = []
        assets: list[dict[str, Any]] = []
        broken_locations: list[str] = []
        incomplete_fields: list[str] = []
        for column in range(1, sheet.max_column + 1):
            label = cell_text(sheet.cell(header_row, column).value)
            if not label:
                continue
            character = cell_text(sheet.cell(1, column).value).upper()
            requirement = cell_text(sheet.cell(2, column).value)
            inferred_asset = asset_type(label)
            if character in ASSET_CHARACTERS or inferred_asset:
                if inferred_asset and inferred_asset not in {item["type"] for item in assets}:
                    assets.append(
                        {"type": inferred_asset, "required": fold(requirement) == "obligatorio"}
                    )
                continue
            if not character or not requirement:
                incomplete_fields.append(sheet.cell(header_row, column).coordinate)
            candidate_key = attribute_key(label)
            identity = candidate_key if candidate_key in BASE_ERP_KEYS else words(label)
            raw_values: list[object] = []
            for row in range(header_row + 1, sheet.max_row + 1):
                value = sheet.cell(row, column).value
                if is_broken_marker(value):
                    broken_locations.append(sheet.cell(row, column).coordinate)
                elif meaningful(value):
                    raw_values.append(value)
            occurrence = {
                "identity": identity,
                "candidateKey": candidate_key,
                "label": label,
                "character": character,
                "receives": cell_text(sheet.cell(3, column).value),
                "unit": explicit_unit(label),
                "rawValues": raw_values,
            }
            occurrences[identity].append(occurrence)
            attributes.append(
                {
                    "identity": identity,
                    "column": column,
                    "required": fold(requirement) == "obligatorio",
                    "replicable": yes(sheet.cell(4, column).value),
                    "includeInTechnicalSheet": yes(sheet.cell(5, column).value),
                    "roles": {
                        "ADMINISTRADOR": yes(sheet.cell(7, column).value),
                        "COMPRAS": yes(sheet.cell(8, column).value),
                        "VENTAS": yes(sheet.cell(9, column).value),
                    },
                }
            )
        if incomplete_fields:
            diagnostics.add(
                "FIELD_METADATA_INCOMPLETE",
                "Hay campos etiquetados sin tipo o condición de obligatoriedad; se conservaron como texto/opcionales según la evidencia disponible.",
                sheet=sheet.title,
                count=len(incomplete_fields),
                details=incomplete_fields,
            )
        if broken_locations:
            diagnostics.add(
                "BROKEN_MARKER_IGNORED",
                "Se ignoraron marcadores de error de Excel en la plantilla.",
                sheet=sheet.title,
                count=len(broken_locations),
                details=broken_locations[:20],
            )

        sku_attribute = next(
            (item for item in attributes if item["identity"] == "codigo_articulo"), None
        )
        if sku_attribute is None:
            raise ValueError(f"{sheet.title}: no contiene Código de Articulo - CDR")
        products: list[dict[str, Any]] = []
        for row in range(header_row + 1, sheet.max_row + 1):
            raw_sku = sheet.cell(row, sku_attribute["column"]).value
            if not meaningful(raw_sku) or fold(raw_sku) == "plano":
                continue
            values = {
                item["identity"]: sheet.cell(row, item["column"]).value
                for item in attributes
                if meaningful(sheet.cell(row, item["column"]).value)
            }
            products.append({"excelRow": row, "rawValues": values})

        raw_templates.append(
            {
                "sheet": sheet.title,
                "slug": template_slug,
                "name": template_name,
                "variantKey": variant_key,
                "categoryNumber": base_number,
                "categorySlug": category["slug"],
                "categoryPath": category["categoryPath"],
                "application": category["application"],
                "headerRow": header_row,
                "attributes": attributes,
                "assets": assets,
                "products": products,
            }
        )

    candidate_identities: dict[str, list[str]] = defaultdict(list)
    for identity, entries in occurrences.items():
        candidate_identities[entries[0]["candidateKey"]].append(identity)
    identity_to_key: dict[str, str] = {}
    for identity, entries in occurrences.items():
        candidate = entries[0]["candidateKey"]
        if len(candidate_identities[candidate]) == 1:
            identity_to_key[identity] = candidate
        else:
            suffix = hashlib.sha1(identity.encode("utf-8")).hexdigest()[:6]
            identity_to_key[identity] = f"{candidate[:93]}_{suffix}"

    definitions: list[dict[str, Any]] = []
    definition_by_identity: dict[str, dict[str, Any]] = {}
    for identity, entries in occurrences.items():
        key = identity_to_key[identity]
        values = [value for entry in entries for value in entry["rawValues"] if meaningful(value)]
        numeric_declared = all(entry["character"].startswith("NUMERICO") for entry in entries)
        all_numeric = bool(values) and all(parse_number(value) is not None for value in values)
        all_boolean = bool(values) and all(parse_boolean(value) is not None for value in values)
        units = {entry["unit"] for entry in entries if entry["unit"]}
        unit = next(iter(units)) if len(units) == 1 and numeric_declared and all_numeric else None
        data_type = (
            "boolean"
            if all_boolean
            else "measurement"
            if unit
            else "number"
            if numeric_declared and all_numeric
            else "text"
        )
        definition = {
            "key": key,
            "label": entries[0]["label"],
            "dataType": data_type,
            "unit": unit,
            "sourceAuthority": source_authority(entries[0]["receives"], key),
        }
        definitions.append(definition)
        definition_by_identity[identity] = definition

    templates: list[dict[str, Any]] = []
    product_count = 0
    for raw_template in raw_templates:
        attributes = []
        for item in raw_template["attributes"]:
            definition = definition_by_identity[item["identity"]]
            roles = {
                role: {
                    "canView": visible,
                    "canEdit": visible and role != "VENTAS" and definition["sourceAuthority"] == "pim",
                    "canImport": visible and role != "VENTAS" and definition["sourceAuthority"] == "pim",
                    "canExport": visible,
                }
                for role, visible in item["roles"].items()
            }
            attributes.append(
                {
                    "key": definition["key"],
                    "required": item["required"],
                    "replicable": item["replicable"],
                    "searchable": True,
                    "includeInTechnicalSheet": item["includeInTechnicalSheet"],
                    "roles": roles,
                }
            )
        products = []
        for raw_product in raw_template["products"]:
            converted: dict[str, object] = {}
            for identity, value in raw_product["rawValues"].items():
                definition = definition_by_identity[identity]
                if definition["dataType"] in {"number", "measurement"}:
                    converted_value = parse_number(value)
                elif definition["dataType"] == "boolean":
                    converted_value = parse_boolean(value)
                else:
                    converted_value = cell_text(value)
                if converted_value is not None and meaningful(converted_value):
                    converted[definition["key"]] = converted_value
            sku = cell_text(converted.get("codigo_articulo"))
            provider = cell_text(converted.get("codigo_proveedor"))
            unified = cell_text(converted.get("codigo_unificador"))
            principal = cell_text(converted.get("categoria_principal"))
            second_level = cell_text(converted.get("categoria_segundo_nivel"))
            third_level = cell_text(converted.get("categoria_tercer_nivel"))
            application = cell_text(converted.get("tipo_aplicacion"))
            brand = cell_text(converted.get("marca"))
            if not sku or not unified:
                raise ValueError(
                    f"{raw_template['sheet']} fila {raw_product['excelRow']}: SKU/código unificador vacío"
                )
            products.append(
                {
                    "excelRow": raw_product["excelRow"],
                    "sku": sku,
                    "providerCode": provider,
                    "unifiedCode": unified,
                    "name": provider or f"{third_level} {sku}".strip(),
                    "categoryPrincipal": principal,
                    "categorySecondLevel": second_level,
                    "categoryThirdLevel": third_level,
                    "application": application,
                    "brand": brand,
                    "values": converted,
                }
            )
        product_count += len(products)
        templates.append(
            {
                key: value
                for key, value in raw_template.items()
                if key not in {"products", "attributes"}
            }
            | {"attributes": attributes, "products": products}
        )

    templates_by_category: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for template in templates:
        templates_by_category[template["categorySlug"]].append(template)
    parent_slugs = {
        category["parentSlug"] for category in categories if category["parentSlug"] is not None
    }
    for category in categories:
        if category["slug"] in parent_slugs:
            category["defined"] = False
            continue
        variants = sorted(
            templates_by_category.get(category["slug"], []),
            key=lambda item: tuple(int(part) for part in item["variantKey"].split(".")),
        )
        category["defined"] = bool(variants)
        category["templateVariants"] = [
            {
                "variantKey": template["variantKey"],
                "templateSheet": template["sheet"],
                "slug": template["slug"],
                "name": template["name"],
            }
            for template in variants
        ]
    unmapped = [
        category
        for category in categories
        if category["slug"] not in parent_slugs and not category["defined"]
    ]
    if unmapped:
        diagnostics.add(
            "UNMAPPED_TAXONOMY_CATEGORIES",
            "La matriz contiene categorías sin plantilla de producto; se conservaron como categorías pendientes.",
            count=len(unmapped),
            details=[category["slug"] for category in unmapped[:20]],
        )

    product_skus = [product["sku"] for template in templates for product in template["products"]]
    duplicate_skus = sorted(sku for sku in set(product_skus) if product_skus.count(sku) > 1)
    if duplicate_skus:
        raise ValueError(f"SKU duplicados en las plantillas: {', '.join(duplicate_skus)}")

    unified_codes = {
        product["unifiedCode"] for template in templates for product in template["products"]
    }
    relations: dict[str, list[dict[str, object]]] = {
        "oem": relation_rows(book["OEM"], diagnostics) if "OEM" in book.sheetnames else [],
        "homologs": relation_rows(book["HOMOLOGOS"], diagnostics)
        if "HOMOLOGOS" in book.sheetnames
        else [],
        "applications": relation_rows(book["APLICACIONES"], diagnostics)
        if "APLICACIONES" in book.sheetnames
        else [],
        "unifiedCodes": relation_rows(book["CODIGO UNIFICADOR"], diagnostics)
        if "CODIGO UNIFICADOR" in book.sheetnames
        else [],
    }
    unknown_relation_codes = sorted(
        {
            cell_text(relation.get("codigo_unificador"))
            for relation_name in ("oem", "homologs", "applications")
            for relation in relations[relation_name]
            if cell_text(relation.get("codigo_unificador")) not in unified_codes
        }
    )
    if unknown_relation_codes:
        raise ValueError(
            "Relaciones con código unificador inexistente: " + ", ".join(unknown_relation_codes)
        )

    return {
        "version": 2,
        "source": {
            "fileName": source_path.name,
            "sha256": source_hash,
            "sheetCount": len(book.sheetnames),
            "templateCount": len(templates),
            "productCount": product_count,
            "diagnosticCount": len(diagnostics.items),
        },
        "diagnostics": diagnostics.items,
        "categories": categories,
        "definitions": sorted(definitions, key=lambda item: item["key"]),
        "templates": templates,
        "relations": relations,
    }


def main() -> None:
    if len(sys.argv) not in {2, 3}:
        raise SystemExit("usage: generate-client-template-manifest.py INPUT.xlsx [OUTPUT.json]")
    source_path = Path(sys.argv[1]).expanduser().resolve()
    output_path = (
        Path(sys.argv[2]).expanduser().resolve()
        if len(sys.argv) == 3
        else Path(__file__).resolve().parents[1]
        / "apps/api/src/database/client-template-manifest.json"
    )
    manifest = build_manifest(source_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(
        f"generated {output_path}: {manifest['source']['templateCount']} templates, "
        f"{len(manifest['definitions'])} definitions, {manifest['source']['productCount']} products, "
        f"{manifest['source']['diagnosticCount']} diagnostics"
    )


if __name__ == "__main__":
    main()
